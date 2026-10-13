from dataclasses import dataclass
from datetime import date
from math import exp, isfinite
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from asset_classes import calculate_asset_class_add_on
from asset_classes.common import AssetClassCalculation
from risk_weights import resolve_counterparty_risk_weight


@dataclass
class CalculatorConfiguration:
    # Trade inputs. Valuation date defaults to start_date when left as None.
    # asset_class: interest_rate, foreign_exchange, credit, equity, commodity.
    asset_class: str = "interest_rate"
    instrument_type: str = "linear"
    transaction_type: str = "standard"
    start_date: date = date(2027, 1, 1)
    end_date: date = date(2028, 1, 1)
    valuation_date: date | None = None
    # Contract notional in its trade currency (IR and credit).
    notional: float = 1_000_000.0
    # Domestic currency per unit of trade currency for IR/credit notional.
    notional_to_domestic_rate: float = 1.0
    collateral: float = 0.0
    mtm: float = 0.0
    # Linear trade direction: +1 long or -1 short in the primary risk factor.
    supervisory_delta: float = 1.0
    # Contract/hedging-set currency (and domestic currency when applicable).
    currency: str = "USD"
    domestic_currency: str = "USD"
    recognized_cva_loss: float = 0.0
    output_path: Path | None = None

    # European vanilla option inputs. Set these when instrument_type is "option".
    option_type: str = "call"
    option_position: str = "bought"
    option_expiry_date: date | None = None
    option_underlying_price: float = 1.0
    option_strike_price: float = 1.0
    option_lambda: float = 0.0

    # FX notionals and spot rates; each rate is domestic currency per 1 unit.
    fx_base_currency: str = "EUR"
    fx_quote_currency: str = "USD"
    fx_base_notional: float = 1_000_000.0
    fx_quote_notional: float = 1_000_000.0
    fx_base_to_domestic_rate: float = 1.0
    fx_quote_to_domestic_rate: float = 1.0

    # Credit, equity, and commodity classifications and underlyings.
    credit_reference_type: str = "single_name"
    credit_rating_bucket: str = "bbb"
    credit_index_grade: str = "ig"
    equity_reference_type: str = "single_name"
    # Equity/commodity unit price; conversion rate maps its currency to domestic.
    underlying_price: float = 100.0
    underlying_units: float = 10_000.0
    underlying_to_domestic_rate: float = 1.0
    commodity_hedging_set: str = "energy"
    commodity_type: str = "other"

    # Counterparty choices: sovereign, pse, mdb, bank, financial,
    # corporate, or individual.
    counterparty_type: str = "bank"
    # Rating bucket: 1=top bucket through 5=lowest; 6=unrated.
    counterparty_rating_bucket: str = "1"
    sovereign_is_canadian: bool = False
    sovereign_assessment_method: str = "rating"
    sovereign_eca_score: str = "1"
    # PSE status: canadian_provincial, other_canadian, or foreign.
    pse_status: str = "canadian_provincial"
    # Foreign PSE risk weight, expressed as a decimal (1.0 = 100%).
    foreign_pse_risk_weight: float = 1.0
    mdb_qualifies_for_zero_weight: bool = True
    financial_institution_treated_as_bank: bool = True
    # Bank treatment: ecra, scra, unrated_100, smsb_simplified,
    # or canadian_dsib.
    bank_treatment: str = "ecra"
    bank_short_term: bool = False
    bank_scra_grade: str = "a"
    bank_meets_enhanced_scra_criteria: bool = False
    bank_simplified_short_term: bool = False
    # Corporate type: rated_general, unrated_general, unrated_sme,
    # us_gse, or specialized.
    corporate_type: str = "rated_general"
    corporate_rating_bucket: str = "1"
    corporate_unrated_treatment: str = "blanket_100"
    # Specialized lending includes rated_1 through rated_5, or the
    # unrated CAR project/object-finance treatment names used below.
    specialized_lending_type: str = "object_or_commodities"


CONFIG = CalculatorConfiguration(
    asset_class="interest_rate",
    instrument_type="linear",
    transaction_type="standard",
    start_date=date(2027, 1, 1),
    end_date=date(2028, 1, 1),
    valuation_date=None,
    notional=1_000_000.0,
    notional_to_domestic_rate=1.0,
    collateral=0.0,
    mtm=0.0,
    supervisory_delta=1.0,
    currency="USD",
    domestic_currency="USD",
    recognized_cva_loss=0.0,
    output_path=None,
    # Set asset-class-specific options below when changing asset_class.
    option_type="call",
    option_position="bought",
    option_expiry_date=None,
    option_underlying_price=1.0,
    option_strike_price=1.0,
    option_lambda=0.0,
    fx_base_currency="EUR",
    fx_quote_currency="USD",
    fx_base_notional=1_000_000.0,
    fx_quote_notional=1_000_000.0,
    fx_base_to_domestic_rate=1.0,
    fx_quote_to_domestic_rate=1.0,
    credit_reference_type="single_name",
    credit_rating_bucket="bbb",
    credit_index_grade="ig",
    equity_reference_type="single_name",
    underlying_price=100.0,
    underlying_units=10_000.0,
    underlying_to_domestic_rate=1.0,
    commodity_hedging_set="energy",
    commodity_type="other",
    counterparty_type="bank",
    counterparty_rating_bucket="1",
    ########################################
    sovereign_is_canadian=False,
    sovereign_assessment_method="rating",
    sovereign_eca_score="1",
    pse_status="canadian_provincial",
    foreign_pse_risk_weight=1.0,
    mdb_qualifies_for_zero_weight=True,
    financial_institution_treated_as_bank=True,
    bank_treatment="ecra",
    bank_short_term=False,
    bank_scra_grade="a",
    bank_meets_enhanced_scra_criteria=False,
    bank_simplified_short_term=False,
    corporate_type="rated_general",
    corporate_rating_bucket="1",
    corporate_unrated_treatment="blanket_100",
    specialized_lending_type="object_or_commodities",
)


@dataclass
class Trade:
    asset_class: str
    instrument_type: str
    transaction_type: str
    start_date: date
    end_date: date
    valuation_date: date
    notional: float
    collateral: float
    mtm: float
    supervisory_delta: float
    currency: str
    domestic_currency: str


@dataclass
class CCRResult:
    asset_class: str
    instrument_type: str
    supervisory_duration: float | None
    maturity_factor: float
    adjusted_notional: float
    effective_notional: float
    supervisory_delta: float
    supervisory_factor: float
    replacement_cost: float
    aggregate_add_on: float
    pfe_multiplier: float
    pfe: float
    ead: float
    recognized_cva_loss: float
    counterparty_ead: float
    risk_weight: float
    rwa: float

class CCRCalculator:
    """Calculate indicative SA-CCR RWA for one unmargined OTC derivative."""

    ALPHA = 1.4
    PFE_MULTIPLIER_FLOOR = 0.05

    def __init__(self, configuration: CalculatorConfiguration):
        """Configure the SA-CCR calculator."""
        self.configuration = configuration

    def calculate_ccr_rwa(
        self,
        trade: Trade,
        asset_class_calculation: AssetClassCalculation,
        risk_weight: float,
        recognized_cva_loss: float = 0.0,
    ) -> CCRResult:
        """Calculate RC, PFE, EAD, and RWA for a single trade."""
        if trade.start_date >= trade.end_date:
            raise ValueError("The end date must be after the start date.")
        if trade.valuation_date >= trade.end_date:
            raise ValueError(
                "The trade has matured on or before the valuation date."
            )
        if trade.asset_class in {"interest_rate", "credit"}:
            if trade.notional < 0:
                raise ValueError("Notional must be non-negative.")
            if not isfinite(trade.notional):
                raise ValueError("Notional must be finite.")
        if not isfinite(trade.mtm) or not isfinite(trade.collateral):
            raise ValueError("MTM and collateral must be finite.")
        if not isfinite(asset_class_calculation.add_on):
            raise ValueError("Asset class add-on must be finite.")
        if not isfinite(risk_weight) or risk_weight < 0:
            raise ValueError("Risk weight must be non-negative.")
        if (
            not isfinite(recognized_cva_loss)
            or recognized_cva_loss < 0
        ):
            raise ValueError("Recognized CVA loss must be non-negative.")

        aggregate_add_on = asset_class_calculation.add_on
        net_mtm = trade.mtm - trade.collateral
        replacement_cost = max(net_mtm, 0.0)

        if aggregate_add_on == 0:
            pfe_multiplier = 1.0
        else:
            pfe_multiplier = min(
                1.0,
                self.PFE_MULTIPLIER_FLOOR
                + (1.0 - self.PFE_MULTIPLIER_FLOOR)
                * exp(
                    net_mtm
                    / (
                        2
                        * (1.0 - self.PFE_MULTIPLIER_FLOOR)
                        * aggregate_add_on
                    )
                ),
            )

        pfe = pfe_multiplier * aggregate_add_on
        ead = self.ALPHA * (replacement_cost + pfe)
        counterparty_ead = max(ead - recognized_cva_loss, 0.0)
        rwa = counterparty_ead * risk_weight

        return CCRResult(
            asset_class=trade.asset_class,
            instrument_type=trade.instrument_type,
            supervisory_duration=asset_class_calculation.supervisory_duration,
            maturity_factor=asset_class_calculation.maturity_factor,
            adjusted_notional=asset_class_calculation.adjusted_notional,
            effective_notional=asset_class_calculation.effective_notional,
            supervisory_delta=asset_class_calculation.supervisory_delta,
            supervisory_factor=asset_class_calculation.supervisory_factor,
            replacement_cost=replacement_cost,
            aggregate_add_on=aggregate_add_on,
            pfe_multiplier=pfe_multiplier,
            pfe=pfe,
            ead=ead,
            recognized_cva_loss=recognized_cva_loss,
            counterparty_ead=counterparty_ead,
            risk_weight=risk_weight,
            rwa=rwa,
        )

    def resolve_counterparty_risk_weight(self) -> tuple[str, float]:
        """Resolve the configured standardized counterparty risk weight."""
        return resolve_counterparty_risk_weight(self.configuration)

    def export_to_excel(
        self,
        trade: Trade,
        result: CCRResult,
        counterparty_name: str,
        asset_class_details: tuple[
            tuple[str, str | float | date], ...
        ],
        output_path: Path | None = None,
    ) -> Path:
        """Write the trade inputs and SA-CCR results to an Excel workbook."""
        if output_path is None:
            output_path = Path(__file__).with_name("CCR_RWA_Calculation.xlsx")

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "CCR RWA"
        worksheet.append(["OSFI SA-CCR RWA Calculation"])
        worksheet.merge_cells("A1:C1")
        worksheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
        worksheet["A1"].fill = PatternFill(
            fill_type="solid", fgColor="1F4E78"
        )
        worksheet["A1"].alignment = Alignment(horizontal="center")
        worksheet.append(["Section", "Metric", "Value"])

        rows = [
            ("Trade inputs", "Asset class", trade.asset_class),
            ("Trade inputs", "Instrument type", trade.instrument_type),
            ("Trade inputs", "Transaction type", trade.transaction_type),
            ("Trade inputs", "Start date", trade.start_date),
            ("Trade inputs", "End date", trade.end_date),
            ("Trade inputs", "Valuation date", trade.valuation_date),
            ("Trade inputs", "Currency", trade.currency),
            ("Trade inputs", "Domestic currency", trade.domestic_currency),
            ("Trade inputs", "Supervisory delta", result.supervisory_delta),
            ("Trade inputs", "MTM", trade.mtm),
            ("Trade inputs", "Net cash collateral", trade.collateral),
            ("Counterparty", "Classification", counterparty_name),
            ("Counterparty", "Risk weight", result.risk_weight),
        ]
        if trade.asset_class in {"interest_rate", "credit"}:
            rows.insert(
                6,
                ("Trade inputs", "Contract notional", trade.notional),
            )
        rows.extend(
            ("Asset class inputs", label, value)
            for label, value in asset_class_details
        )
        if result.supervisory_duration is not None:
            rows.append(
                (
                    "SA-CCR",
                    "Supervisory duration (years)",
                    result.supervisory_duration,
                )
            )
        rows.extend([
            ("SA-CCR", "Maturity factor", result.maturity_factor),
            ("SA-CCR", "Adjusted notional", result.adjusted_notional),
            ("SA-CCR", "Effective notional", result.effective_notional),
            ("SA-CCR", "Supervisory factor", result.supervisory_factor),
            ("SA-CCR", "Replacement cost (RC)", result.replacement_cost),
            ("SA-CCR", "Aggregate add-on", result.aggregate_add_on),
            ("SA-CCR", "PFE multiplier", result.pfe_multiplier),
            ("SA-CCR", "PFE", result.pfe),
            ("SA-CCR", "EAD before CVA loss adjustment", result.ead),
            ("SA-CCR", "Recognized CVA loss", result.recognized_cva_loss),
            ("SA-CCR", "EAD after CVA loss adjustment", result.counterparty_ead),
            ("SA-CCR", "CCR RWA", result.rwa),
        ])

        for row in rows:
            worksheet.append(row)

        for cell in worksheet[2]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="5B9BD5")

        for row_number in range(3, worksheet.max_row + 1):
            worksheet.cell(row_number, 1).font = Font(bold=True)
            metric = worksheet.cell(row_number, 2).value
            value_cell = worksheet.cell(row_number, 3)
            if metric in {"Risk weight", "Supervisory factor"}:
                value_cell.number_format = "0.00%"
            elif metric in {"Supervisory delta", "PFE multiplier"}:
                value_cell.number_format = "0.000000"
            elif isinstance(value_cell.value, (float, int)):
                value_cell.number_format = "#,##0.00"
            if isinstance(value_cell.value, date):
                value_cell.number_format = "yyyy-mm-dd"

        worksheet.column_dimensions["A"].width = 18
        worksheet.column_dimensions["B"].width = 40
        worksheet.column_dimensions["C"].width = 30
        worksheet.freeze_panes = "A3"
        worksheet.auto_filter.ref = f"A2:C{worksheet.max_row}"

        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
        return output_path

    def run(self) -> CCRResult:
        print("OSFI SA-CCR RWA calculator")
        print(
            "Scope: one unmargined, non-CCP OTC trade; "
            "linear products and European vanilla options."
        )
        config = self.configuration
        if not config.currency or not config.domestic_currency:
            raise ValueError("Trade and domestic currencies must not be empty.")
        if config.start_date >= config.end_date:
            raise ValueError("The end date must be after the start date.")
        if (config.valuation_date or config.start_date) >= config.end_date:
            raise ValueError(
                "The trade has matured on or before the valuation date."
            )

        asset_class_calculation = calculate_asset_class_add_on(config)
        trade = Trade(
            asset_class=config.asset_class,
            instrument_type=config.instrument_type,
            transaction_type=config.transaction_type,
            start_date=config.start_date,
            end_date=config.end_date,
            valuation_date=config.valuation_date or config.start_date,
            notional=config.notional,
            collateral=config.collateral,
            mtm=config.mtm,
            supervisory_delta=config.supervisory_delta,
            currency=config.currency.upper(),
            domestic_currency=config.domestic_currency.upper(),
        )
        counterparty_name, risk_weight = (
            self.resolve_counterparty_risk_weight()
        )
        output_path = config.output_path

        result = self.calculate_ccr_rwa(
            trade,
            asset_class_calculation,
            risk_weight,
            config.recognized_cva_loss,
        )
        output_path = self.export_to_excel(
            trade,
            result,
            counterparty_name,
            asset_class_calculation.details,
            output_path,
        )
        print(f"Excel output saved to: {output_path}")
        return result


if __name__ == "__main__":
    calculator = CCRCalculator(CONFIG)
    calculator.run()
